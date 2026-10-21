"""Поведенческие тесты core/github_sync.py и core/pptx_renderer.py.

`github_sync` — обёртки над git: проверяем разбор вывода и обработку сбоев
на подменённом `subprocess.run` (реальный репозиторий не трогаем).
`pptx_renderer` — разбор HTML на слайды и настоящая сборка .pptx через
python-pptx во временном каталоге.
"""
import os
import sys
import tempfile
import types

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.core import github_sync as gs
from src.core import pptx_renderer as pr
from src.core.pptx_renderer import PPTXRenderer, _SlideParser


# ─── фальшивый subprocess ────────────────────────────────────────────────────

class _Result:
    def __init__(self, stdout="", stderr="", returncode=0):
        self.stdout = stdout
        self.stderr = stderr
        self.returncode = returncode


class _FakeSubprocess:
    """Подменяет `subprocess` в github_sync, записывая команды."""

    def __init__(self, responses=None):
        self.calls = []
        self.responses = list(responses or [])

    def run(self, cmd, **kw):
        self.calls.append((list(cmd), kw))
        for expected, response in self.responses:
            if list(cmd)[:len(expected)] == list(expected):
                if isinstance(response, Exception):
                    raise response
                return response
        return _Result()

    def commands(self):
        return [c for c, _ in self.calls]


class _PatchedGit:
    def __init__(self, responses=None):
        self.fake = _FakeSubprocess(responses)
        self._orig = gs.subprocess

    def __enter__(self):
        gs.subprocess = types.SimpleNamespace(run=self.fake.run)
        return self.fake

    def __exit__(self, *exc):
        gs.subprocess = self._orig
        return False


# ─── github_sync: базовые операции ───────────────────────────────────────────

def test_is_git_repo_true():
    """Нулевой код возврата означает «это git-репозиторий»."""
    with _PatchedGit([(["git", "rev-parse"], _Result(returncode=0))]) as fake:
        assert gs.is_git_repo("repo") is True
    cmd, kw = fake.calls[0]
    assert cmd == ["git", "rev-parse", "--git-dir"]
    assert kw["cwd"] == "repo"
    assert kw["timeout"] == 5
    print("[PASS] test_is_git_repo_true")


def test_is_git_repo_false_on_nonzero_exit():
    """Ненулевой код — не репозиторий."""
    with _PatchedGit([(["git", "rev-parse"], _Result(returncode=128))]):
        assert gs.is_git_repo("repo") is False
    print("[PASS] test_is_git_repo_false_on_nonzero_exit")


def test_is_git_repo_false_on_exception():
    """Исключение запуска git не выходит наружу."""
    with _PatchedGit([(["git", "rev-parse"], OSError("git not found"))]):
        assert gs.is_git_repo("repo") is False
    print("[PASS] test_is_git_repo_false_on_exception")


def test_git_status_classifies_porcelain_lines():
    """Строки porcelain раскладываются по staged/modified/untracked."""
    out = "\n".join([
        "M  staged.py",
        "?? new.py",
        "R  old.py -> renamed.py",
        " D deleted.py",
    ])
    with _PatchedGit([(["git", "status"], _Result(stdout=out))]):
        status = gs.git_status("repo")
    assert status["staged"] == ["staged.py", "deleted.py"], status
    assert status["untracked"] == ["new.py"], status
    assert status["modified"] == ["old.py -> renamed.py"], status
    assert status["raw"] == out.splitlines()
    print("[PASS] test_git_status_classifies_porcelain_lines")


def test_git_status_empty_output():
    """Пустой вывод — пустые списки."""
    with _PatchedGit([(["git", "status"], _Result(stdout="   \n"))]):
        status = gs.git_status("repo")
    assert status == {"staged": [], "modified": [], "untracked": [], "raw": []}, status
    print("[PASS] test_git_status_empty_output")


def test_git_status_error_is_captured():
    """Сбой git status возвращает безопасный словарь с текстом ошибки."""
    with _PatchedGit([(["git", "status"], OSError("boom"))]):
        status = gs.git_status("repo")
    assert status["staged"] == [] and status["raw"] == []
    assert status["error"] == "boom", status
    print("[PASS] test_git_status_error_is_captured")


def test_git_commit_stages_given_files():
    """Указанные файлы добавляются по одному, затем коммит."""
    with _PatchedGit([(["git", "commit"], _Result(stdout="done\n", stderr=""))]) as fake:
        result = gs.git_commit("repo", "msg", ["a.py", "b.py"])
    assert result == {"success": True, "output": "done", "error": ""}, result
    assert fake.commands() == [
        ["git", "add", "a.py"],
        ["git", "add", "b.py"],
        ["git", "commit", "-m", "msg"],
    ], fake.commands()
    print("[PASS] test_git_commit_stages_given_files")


def test_git_commit_stages_all_without_file_list():
    """Без списка файлов используется git add -A."""
    with _PatchedGit() as fake:
        gs.git_commit("repo", "msg")
    assert fake.commands()[0] == ["git", "add", "-A"], fake.commands()
    print("[PASS] test_git_commit_stages_all_without_file_list")


def test_git_commit_reports_failure():
    """Ненулевой код коммита — success False с текстом stderr."""
    with _PatchedGit([(["git", "commit"], _Result(returncode=1, stderr="nothing to commit"))]):
        result = gs.git_commit("repo", "msg")
    assert result["success"] is False
    assert result["error"] == "nothing to commit"
    print("[PASS] test_git_commit_reports_failure")


def test_git_commit_exception_is_caught():
    """Исключение при коммите не выходит наружу."""
    with _PatchedGit([(["git", "add"], OSError("locked"))]):
        result = gs.git_commit("repo", "msg")
    assert result == {"success": False, "error": "locked"}, result
    print("[PASS] test_git_commit_exception_is_caught")


def test_git_push_detects_current_branch():
    """Без branch выполняется git branch --show-current, затем push."""
    with _PatchedGit([
        (["git", "branch"], _Result(stdout="dev\n")),
        (["git", "push"], _Result(stdout="ok\n")),
    ]) as fake:
        result = gs.git_push("repo")
    assert result["success"] is True
    assert fake.commands() == [["git", "branch", "--show-current"],
                               ["git", "push", "origin", "dev"]], fake.commands()
    print("[PASS] test_git_push_detects_current_branch")


def test_git_push_falls_back_to_main():
    """Пустой вывод текущей ветки трактуется как main."""
    with _PatchedGit([(["git", "branch"], _Result(stdout="\n"))]) as fake:
        gs.git_push("repo")
    assert fake.commands()[-1] == ["git", "push", "origin", "main"], fake.commands()
    print("[PASS] test_git_push_falls_back_to_main")


def test_git_push_with_explicit_branch_skips_detection():
    """Явная ветка избавляет от лишнего вызова git branch."""
    with _PatchedGit() as fake:
        result = gs.git_push("repo", remote="upstream", branch="feature")
    assert fake.commands() == [["git", "push", "upstream", "feature"]], fake.commands()
    assert result["success"] is True
    print("[PASS] test_git_push_with_explicit_branch_skips_detection")


def test_git_push_reports_error():
    """Ошибка push отдаётся в поле error."""
    with _PatchedGit([(["git", "push"], _Result(returncode=1, stderr="rejected"))]):
        result = gs.git_push("repo", branch="main")
    assert result == {"success": False, "output": "", "error": "rejected"}, result
    print("[PASS] test_git_push_reports_error")


def test_git_push_exception_is_caught():
    """Исключение при push не выходит наружу."""
    with _PatchedGit([(["git", "push"], OSError("offline"))]):
        result = gs.git_push("repo", branch="main")
    assert result == {"success": False, "error": "offline"}, result
    print("[PASS] test_git_push_exception_is_caught")


def test_git_pull_detects_branch_and_returns_output():
    """Pull сам определяет ветку и возвращает её вывод."""
    with _PatchedGit([
        (["git", "branch"], _Result(stdout="main\n")),
        (["git", "pull"], _Result(stdout="up to date\n")),
    ]) as fake:
        result = gs.git_pull("repo")
    assert result == {"success": True, "output": "up to date", "error": ""}, result
    assert fake.commands() == [["git", "branch", "--show-current"],
                               ["git", "pull", "origin", "main"]], fake.commands()
    print("[PASS] test_git_pull_detects_branch_and_returns_output")


def test_git_pull_explicit_branch_and_exception():
    """Явная ветка и сбой pull обрабатываются корректно."""
    with _PatchedGit() as fake:
        gs.git_pull("repo", branch="dev")
    assert fake.commands() == [["git", "pull", "origin", "dev"]], fake.commands()

    with _PatchedGit([(["git", "pull"], OSError("no network"))]):
        result = gs.git_pull("repo", branch="dev")
    assert result == {"success": False, "error": "no network"}, result
    print("[PASS] test_git_pull_explicit_branch_and_exception")


def test_git_log_parses_and_shortens_hash():
    """Коммиты разбираются, хеш укорачивается до 8 символов."""
    stdout = "\n".join([
        "abcdef0123456789|Первый коммит|2026-10-21 10:00:00 +0300|Иван",
        "0123456789abcdef|Второй коммит|2026-10-20 09:00:00 +0300|Пётр",
        "мусор без разделителей",
    ])
    with _PatchedGit([(["git", "log"], _Result(stdout=stdout))]) as fake:
        commits = gs.git_log("repo", count=7)
    assert len(commits) == 2, commits
    assert commits[0] == {
        "hash": "abcdef01",
        "message": "Первый коммит",
        "date": "2026-10-21 10:00:00 +0300",
        "author": "Иван",
    }, commits[0]
    assert commits[1]["hash"] == "01234567"
    assert commits[1]["author"] == "Пётр"
    assert "--max-count=7" in fake.commands()[0], fake.commands()
    print("[PASS] test_git_log_parses_and_shortens_hash")


def test_git_log_keeps_extra_pipes_in_last_field():
    """split('|', 3) — дополнительные пайпы остаются в поле автора."""
    stdout = "abcdef0123456789|msg|date|автор|лишнее"
    with _PatchedGit([(["git", "log"], _Result(stdout=stdout))]):
        commits = gs.git_log("repo")
    assert commits[0]["author"] == "автор|лишнее", commits[0]
    print("[PASS] test_git_log_keeps_extra_pipes_in_last_field")


def test_git_log_error_returns_empty_list():
    """Сбой git log даёт пустой список."""
    with _PatchedGit([(["git", "log"], OSError("bad repo"))]):
        assert gs.git_log("repo") == []
    print("[PASS] test_git_log_error_returns_empty_list")


def test_git_init_success_and_failure():
    """git init возвращает успех/ошибку по коду и исключению."""
    with _PatchedGit([(["git", "init"], _Result(stdout="Initialized\n"))]):
        assert gs.git_init("repo") == {"success": True, "output": "Initialized"}
    with _PatchedGit([(["git", "init"], _Result(returncode=1))]):
        assert gs.git_init("repo")["success"] is False
    with _PatchedGit([(["git", "init"], OSError("denied"))]):
        assert gs.git_init("repo") == {"success": False, "error": "denied"}
    print("[PASS] test_git_init_success_and_failure")


# ─── github_sync: расширенные операции ───────────────────────────────────────

def test_git_diff_variants():
    """git_diff собирает команду в зависимости от staged/file_path."""
    with _PatchedGit([(["git", "diff"], _Result(stdout="diff --git", stderr="warn"))]) as fake:
        result = gs.git_diff("repo")
    assert result == {"success": True, "diff": "diff --git", "error": "warn"}, result
    assert fake.commands()[0] == ["git", "diff"], fake.commands()

    with _PatchedGit() as fake:
        gs.git_diff("repo", file_path="a.py", staged=True)
    assert fake.commands()[0] == ["git", "diff", "--staged", "a.py"], fake.commands()

    with _PatchedGit([(["git", "diff"], OSError("broken"))]):
        result = gs.git_diff("repo")
    assert result == {"success": False, "diff": "", "error": "broken"}, result
    print("[PASS] test_git_diff_variants")


def test_git_diff_commits():
    """git_diff_commits передаёт оба коммита и возвращает diff."""
    with _PatchedGit([(["git", "diff"], _Result(stdout="delta"))]) as fake:
        result = gs.git_diff_commits("repo", "aaa", "bbb")
    assert result == {"success": True, "diff": "delta", "error": ""}, result
    assert fake.commands()[0] == ["git", "diff", "aaa", "bbb"], fake.commands()

    with _PatchedGit([(["git", "diff"], OSError("nope"))]):
        assert gs.git_diff_commits("repo", "a", "b")["success"] is False
    print("[PASS] test_git_diff_commits")


def test_git_branch_list_marks_current_and_strips_remotes():
    """Текущая ветка помечается, remotes/ из имён убирается."""
    stdout = "\n".join([
        "* main",
        "  feature/x",
        "  remotes/origin/main",
        "   ",
    ])
    with _PatchedGit([(["git", "branch"], _Result(stdout=stdout))]):
        info = gs.git_branch_list("repo")
    assert info["success"] is True
    assert info["current"] == "main", info
    assert info["branches"] == ["main", "feature/x", "origin/main"], info
    print("[PASS] test_git_branch_list_marks_current_and_strips_remotes")


def test_git_branch_list_error():
    """Сбой списка веток — success False и пустые значения."""
    with _PatchedGit([(["git", "branch"], OSError("bad"))]):
        info = gs.git_branch_list("repo")
    assert info == {"success": False, "branches": [], "current": "", "error": "bad"}, info
    print("[PASS] test_git_branch_list_error")


def test_git_branch_create_and_switch():
    """Создание и переключение ветки передают имя."""
    with _PatchedGit([(["git", "branch"], _Result(stdout=""))]) as fake:
        assert gs.git_branch_create("repo", "new")["success"] is True
    assert fake.commands()[0] == ["git", "branch", "new"], fake.commands()

    with _PatchedGit([(["git", "checkout"], _Result(returncode=1, stderr="exists"))]):
        result = gs.git_branch_switch("repo", "new")
    assert result["success"] is False and result["error"] == "exists", result

    with _PatchedGit([(["git", "checkout"], OSError("x"))]):
        assert gs.git_branch_switch("repo", "new")["success"] is False
    with _PatchedGit([(["git", "branch"], OSError("x"))]):
        assert gs.git_branch_create("repo", "new")["success"] is False
    print("[PASS] test_git_branch_create_and_switch")


def test_git_stash_and_pop():
    """Stash и stash pop возвращают успех и вывод."""
    with _PatchedGit([(["git", "stash", "pop"], _Result(stdout="popped"))]) as fake:
        result = gs.git_stash_pop("repo")
    assert result == {"success": True, "output": "popped", "error": ""}, result
    assert fake.commands()[0] == ["git", "stash", "pop"], fake.commands()

    with _PatchedGit([(["git", "stash"], _Result(stdout="saved"))]):
        assert gs.git_stash("repo")["output"] == "saved"

    with _PatchedGit([(["git", "stash"], OSError("x"))]):
        assert gs.git_stash("repo")["success"] is False
    with _PatchedGit([(["git", "stash"], OSError("x"))]):
        assert gs.git_stash_pop("repo")["success"] is False
    print("[PASS] test_git_stash_and_pop")


def test_git_blame_returns_porcelain():
    """Blame отдаёт сырой вывод porcelain."""
    with _PatchedGit([(["git", "blame"], _Result(stdout="abc 1 1 1", stderr=""))]) as fake:
        result = gs.git_blame("repo", "a.py")
    assert result == {"success": True, "blame": "abc 1 1 1", "error": ""}, result
    assert fake.commands()[0] == ["git", "blame", "--porcelain", "a.py"], fake.commands()

    with _PatchedGit([(["git", "blame"], OSError("no file"))]):
        result = gs.git_blame("repo", "a.py")
    assert result == {"success": False, "blame": "", "error": "no file"}, result
    print("[PASS] test_git_blame_returns_porcelain")


def test_git_tag_with_and_without_message():
    """Тег создаётся аннотированным; -m добавляется только при сообщении."""
    with _PatchedGit() as fake:
        gs.git_tag("repo", "v1.0", "Релиз")
    assert fake.commands()[0] == ["git", "tag", "-a", "v1.0", "-m", "Релиз"], fake.commands()

    with _PatchedGit() as fake:
        gs.git_tag("repo", "v2.0")
    assert fake.commands()[0] == ["git", "tag", "-a", "v2.0"], fake.commands()

    with _PatchedGit([(["git", "tag"], OSError("x"))]):
        assert gs.git_tag("repo", "v3")["success"] is False
    print("[PASS] test_git_tag_with_and_without_message")


def test_git_tags_list():
    """Список тегов — по строке на тег; пустой вывод даёт []."""
    with _PatchedGit([(["git", "tag"], _Result(stdout="v1.0\nv1.1\n"))]):
        assert gs.git_tags("repo") == ["v1.0", "v1.1"]
    with _PatchedGit([(["git", "tag"], _Result(stdout=""))]):
        assert gs.git_tags("repo") == []
    with _PatchedGit([(["git", "tag"], OSError("x"))]):
        assert gs.git_tags("repo") == []
    print("[PASS] test_git_tags_list")


def test_git_file_history_filters_by_path():
    """История файла ограничивается путём и числом коммитов."""
    stdout = "abcdef0123456789|Правка|2026-10-21 10:00:00 +0300|Иван"
    with _PatchedGit([(["git", "log"], _Result(stdout=stdout))]) as fake:
        history = gs.git_file_history("repo", "src/a.py", count=4)
    assert history == [{
        "hash": "abcdef01",
        "message": "Правка",
        "date": "2026-10-21 10:00:00 +0300",
        "author": "Иван",
    }], history
    cmd = fake.commands()[0]
    assert "--max-count=4" in cmd and cmd[-2:] == ["--", "src/a.py"], cmd

    with _PatchedGit([(["git", "log"], OSError("x"))]):
        assert gs.git_file_history("repo", "a.py") == []
    print("[PASS] test_git_file_history_filters_by_path")


# ─── pptx_renderer: разбор HTML ──────────────────────────────────────────────

def _parsed(html):
    """Прогоняет HTML через парсер и сбрасывает хвост в последний слайд."""
    parser = _SlideParser()
    parser.feed(html)
    parser.close()
    parser._flush_content()
    return parser


def test_slide_parser_splits_on_headings():
    """Каждый h1/h2 начинает новый слайд."""
    parser = _parsed("<h1>Первый</h1><p>текст один</p><h2>Второй</h2><p>текст два</p>")
    titles = [s["title"] for s in parser.slides]
    assert titles == ["Первый", "Второй"], titles
    assert parser.slides[0]["content"] == ["текст один"], parser.slides
    assert parser.slides[1]["content"] == ["текст два"], parser.slides
    print("[PASS] test_slide_parser_splits_on_headings")


def test_slide_parser_content_before_first_heading():
    """Контент до первого заголовка даёт вводный слайд без названия."""
    parser = _parsed("<p>вступление</p><h1>Основное</h1><p>тело</p>")
    assert parser.slides[0]["title"] == "", parser.slides
    assert parser.slides[0]["content"] == ["вступление"], parser.slides
    assert parser.slides[1]["content"] == ["тело"], parser.slides
    print("[PASS] test_slide_parser_content_before_first_heading")


def test_slide_parser_skips_script_and_style():
    """Содержимое script/style не попадает в слайды."""
    parser = _parsed("<h1>Слайд</h1><script>alert(1)</script><style>.a{}</style><p>видно</p>")
    assert parser.slides[0]["content"] == ["видно"], parser.slides
    print("[PASS] test_slide_parser_skips_script_and_style")


def test_slide_parser_entities_in_heading():
    """HTML-сущности в заголовке декодируются парсером."""
    parser = _parsed("<h1>A &amp; B &#169;</h1>")
    assert parser.slides[0]["title"] == "A & B ©", parser.slides
    print("[PASS] test_slide_parser_entities_in_heading")


def test_slide_parser_entity_handlers_keep_raw_form():
    """Ручные обработчики сущностей сохраняют исходную запись."""
    parser = _SlideParser()
    parser._in_heading = True
    parser._current_title = "A "
    parser.handle_entityref("amp")
    parser.handle_charref("169")
    assert parser._current_title == "A &amp;&#169;", parser._current_title

    # Вне заголовка (current_title = None) обработчики ничего не делают
    parser._current_title = None
    parser.handle_entityref("amp")
    parser.handle_charref("169")
    assert parser._current_title is None
    print("[PASS] test_slide_parser_entity_handlers_keep_raw_form")


def test_slide_parser_joins_text_within_paragraph():
    """Кусочки текста внутри абзаца склеиваются через пробел."""
    parser = _parsed("<h1>Т</h1><p><b>жирный</b> обычный</p>")
    assert parser.slides[0]["content"] == ["жирный обычный"], parser.slides
    print("[PASS] test_slide_parser_joins_text_within_paragraph")


def test_slide_parser_empty_heading_is_ignored():
    """Пустой заголовок не создаёт слайд."""
    parser = _parsed("<h1>   </h1><p>текст</p>")
    assert "".join(s["title"] for s in parser.slides) == "", parser.slides
    assert parser.slides[0]["content"] == ["текст"], parser.slides
    print("[PASS] test_slide_parser_empty_heading_is_ignored")


# ─── pptx_renderer: слайды ───────────────────────────────────────────────────

def test_parse_slides_adds_title_slide_with_author():
    """Титульный слайд содержит заголовок и автора."""
    slides = PPTXRenderer()._parse_slides("<h1>Раздел</h1>", {"title": "Книга", "author": "Иван"})
    assert slides[0] == {"title": "Книга", "content": ["by Иван"]}, slides
    assert slides[1]["title"] == "Раздел"
    print("[PASS] test_parse_slides_adds_title_slide_with_author")


def test_parse_slides_title_without_author():
    """Без автора титульный слайд пуст по содержимому."""
    slides = PPTXRenderer()._parse_slides("<div>текст</div>", {"title": "Книга"})
    assert slides[0] == {"title": "Книга", "content": []}, slides
    assert slides[1] == {"title": "Книга", "content": ["текст"]}, slides
    print("[PASS] test_parse_slides_title_without_author")


def test_parse_slides_keeps_last_slide_content():
    """Тело финального слайда не теряется (регрессия BUG_009)."""
    slides = PPTXRenderer()._parse_slides(
        "<h1>Первый</h1><p>один</p><h1>Второй</h1><p>два</p>", {})
    assert slides[0]["content"] == ["один"], slides
    assert slides[1]["content"] == ["два"], slides
    print("[PASS] test_parse_slides_keeps_last_slide_content")


def test_parse_slides_fallback_plain_text():
    """Без заголовков создаётся один слайд из очищенного HTML."""
    slides = PPTXRenderer()._parse_slides("<div>первая строка</div><div>вторая</div>", {})
    assert len(slides) == 1, slides
    assert slides[0]["title"] == "Presentation"
    assert slides[0]["content"] == ["первая строка вторая"], slides[0]
    print("[PASS] test_parse_slides_fallback_plain_text")


def test_parse_slides_fallback_caps_content_at_20_lines():
    """Резервный слайд содержит не более 20 строк."""
    body = "".join(f"<div>строка {i}</div>" for i in range(30))
    slides = PPTXRenderer()._parse_slides(body, {})
    assert len(slides[0]["content"]) <= 20, len(slides[0]["content"])
    print("[PASS] test_parse_slides_fallback_caps_content_at_20_lines")


def test_parse_slides_uses_metadata_title_in_fallback():
    """В резервном слайде берётся заголовок из метаданных."""
    slides = PPTXRenderer()._parse_slides("<div>текст</div>", {"title": "Отчёт"})
    assert len(slides) == 2, slides
    assert slides[0]["title"] == "Отчёт"
    assert slides[1]["title"] == "Отчёт", slides
    print("[PASS] test_parse_slides_uses_metadata_title_in_fallback")


def test_slide_parser_ignores_tags_inside_heading():
    """Вложенные теги внутри заголовка не сбивают разбор."""
    parser = _parsed("<h1>За<b>голо</b>вок</h1><p>тело</p>")
    assert parser.slides[0]["title"] == "Заголовок", parser.slides
    assert parser.slides[0]["content"] == ["тело"], parser.slides
    print("[PASS] test_slide_parser_ignores_tags_inside_heading")


def test_render_skips_blank_content_lines():
    """Пустые строки тела слайда пропускаются рендерером."""
    class _BlankContentRenderer(PPTXRenderer):
        def _parse_slides(self, html_content, metadata):
            return [{"title": "Слайд", "content": ["", "   ", "настоящая строка"]}]

    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "deck.pptx")
        count = _BlankContentRenderer().render("", {}, out)
        assert count == 1
        from pptx import Presentation
        prs = Presentation(out)
        assert _texts(prs.slides[0]) == ["настоящая строка"], _texts(prs.slides[0])
    print("[PASS] test_render_skips_blank_content_lines")


# ─── pptx_renderer: сборка файла ─────────────────────────────────────────────

def _texts(slide):
    """Все непустые тексты абзацев тела слайда."""
    return [p.text for p in slide.placeholders[1].text_frame.paragraphs if p.text]


def test_render_creates_file_and_returns_slide_count():
    """render сохраняет валидный .pptx и возвращает число слайдов."""
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "deck.pptx")
        count = PPTXRenderer().render(
            "<h1>Первый</h1><p>один</p><h1>Второй</h1><p>два</p>",
            {"title": "Колода", "author": "Аня"}, out,
        )
        assert count == 3, count
        assert os.path.isfile(out)

        from pptx import Presentation
        prs = Presentation(out)
        assert len(prs.slides) == 3
        assert prs.slides[0].shapes.title.text == "Колода"
        assert prs.slides[1].shapes.title.text == "Первый"
        assert _texts(prs.slides[1]) == ["один"], _texts(prs.slides[1])
    print("[PASS] test_render_creates_file_and_returns_slide_count")


def test_render_uses_default_16_9_size():
    """По умолчанию слайд 13.333×7.5 дюйма."""
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "deck.pptx")
        PPTXRenderer().render("<h1>A</h1>", {}, out)
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(out)
        assert prs.slide_width == Inches(13.333)
        assert prs.slide_height == Inches(7.5)
    print("[PASS] test_render_uses_default_16_9_size")


def test_render_applies_custom_dimensions():
    """Заданные размеры применяются как EMU."""
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "deck.pptx")
        PPTXRenderer().render("<h1>A</h1>", {}, out, slide_width=9144000, slide_height=6858000)
        from pptx import Presentation
        prs = Presentation(out)
        assert prs.slide_width == 9144000
        assert prs.slide_height == 6858000
    print("[PASS] test_render_applies_custom_dimensions")


def test_render_converts_markdown_bullets():
    """Строки «- », «* », «+ » превращаются в буллеты."""
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "deck.pptx")
        PPTXRenderer().render("<h1>Список</h1><p>- первый</p>", {}, out)
        from pptx import Presentation
        prs = Presentation(out)
        assert _texts(prs.slides[0]) == ["• первый"], _texts(prs.slides[0])
    print("[PASS] test_render_converts_markdown_bullets")


def test_render_keeps_numbered_lines():
    """Нумерованные строки остаются как есть."""
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "deck.pptx")
        PPTXRenderer().render("<h1>План</h1><p>1. первый шаг</p>", {}, out)
        from pptx import Presentation
        prs = Presentation(out)
        assert _texts(prs.slides[0]) == ["1. первый шаг"], _texts(prs.slides[0])
    print("[PASS] test_render_keeps_numbered_lines")


def test_render_emphasises_hash_lines():
    """Строка, начинающаяся с #, становится жирной и крупнее."""
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "deck.pptx")
        PPTXRenderer().render("<h1>Раздел</h1><p># подзаголовок</p>", {}, out)
        from pptx import Presentation
        from pptx.util import Pt
        prs = Presentation(out)
        para = [p for p in prs.slides[0].placeholders[1].text_frame.paragraphs if p.text][0]
        assert para.text == "# подзаголовок"
        assert para.font.bold is True
        assert para.font.size == Pt(24)
    print("[PASS] test_render_emphasises_hash_lines")


def test_render_creates_missing_output_directory():
    """Каталог назначения создаётся автоматически."""
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "nested", "deeper", "deck.pptx")
        PPTXRenderer().render("<h1>A</h1>", {}, out)
        assert os.path.isfile(out)
    print("[PASS] test_render_creates_missing_output_directory")


def test_render_without_python_pptx_raises_importerror():
    """Без python-pptx рендерер сообщает, что установить."""
    previous = sys.modules.get("pptx", "missing")
    sys.modules["pptx"] = None
    try:
        try:
            PPTXRenderer().render("<h1>A</h1>", {}, "whatever.pptx")
        except ImportError as exc:
            assert "python-pptx" in str(exc), exc
        else:
            raise AssertionError("ожидался ImportError")
    finally:
        if previous == "missing":
            sys.modules.pop("pptx", None)
        else:
            sys.modules["pptx"] = previous
    print("[PASS] test_render_without_python_pptx_raises_importerror")


def main():
    tests = [
        test_is_git_repo_true,
        test_is_git_repo_false_on_nonzero_exit,
        test_is_git_repo_false_on_exception,
        test_git_status_classifies_porcelain_lines,
        test_git_status_empty_output,
        test_git_status_error_is_captured,
        test_git_commit_stages_given_files,
        test_git_commit_stages_all_without_file_list,
        test_git_commit_reports_failure,
        test_git_commit_exception_is_caught,
        test_git_push_detects_current_branch,
        test_git_push_falls_back_to_main,
        test_git_push_with_explicit_branch_skips_detection,
        test_git_push_reports_error,
        test_git_push_exception_is_caught,
        test_git_pull_detects_branch_and_returns_output,
        test_git_pull_explicit_branch_and_exception,
        test_git_log_parses_and_shortens_hash,
        test_git_log_keeps_extra_pipes_in_last_field,
        test_git_log_error_returns_empty_list,
        test_git_init_success_and_failure,
        test_git_diff_variants,
        test_git_diff_commits,
        test_git_branch_list_marks_current_and_strips_remotes,
        test_git_branch_list_error,
        test_git_branch_create_and_switch,
        test_git_stash_and_pop,
        test_git_blame_returns_porcelain,
        test_git_tag_with_and_without_message,
        test_git_tags_list,
        test_git_file_history_filters_by_path,
        test_slide_parser_splits_on_headings,
        test_slide_parser_content_before_first_heading,
        test_slide_parser_skips_script_and_style,
        test_slide_parser_entities_in_heading,
        test_slide_parser_entity_handlers_keep_raw_form,
        test_slide_parser_joins_text_within_paragraph,
        test_slide_parser_empty_heading_is_ignored,
        test_slide_parser_ignores_tags_inside_heading,
        test_render_skips_blank_content_lines,
        test_parse_slides_adds_title_slide_with_author,
        test_parse_slides_title_without_author,
        test_parse_slides_keeps_last_slide_content,
        test_parse_slides_fallback_plain_text,
        test_parse_slides_fallback_caps_content_at_20_lines,
        test_parse_slides_uses_metadata_title_in_fallback,
        test_render_creates_file_and_returns_slide_count,
        test_render_uses_default_16_9_size,
        test_render_applies_custom_dimensions,
        test_render_converts_markdown_bullets,
        test_render_keeps_numbered_lines,
        test_render_emphasises_hash_lines,
        test_render_creates_missing_output_directory,
        test_render_without_python_pptx_raises_importerror,
    ]
    for t in tests:
        t()
    print("\nAll github_sync/pptx_renderer tests passed!")


if __name__ == "__main__":
    main()
