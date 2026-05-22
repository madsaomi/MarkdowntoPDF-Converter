import os
import logging
from docx import Document
from htmldocx import HtmlToDocx
from typing import Dict, Any, Optional

logger = logging.getLogger("MarkdownToPDF.docx_renderer")


class DOCXRenderer:
    def __init__(self) -> None:
        pass

    def render(self, html_content: str, metadata: Dict[str, Any], output_path: str, reference_docx: Optional[str] = None, lang: str = "ru") -> None:
        """
        Takes raw HTML content (from Markdown parser), converts it to Word format,
        and saves it to the output_path.
        """
        # Create a new Word Document
        if reference_docx and os.path.exists(reference_docx):
            try:
                doc = Document(reference_docx)
            except Exception as e:
                logger.warning("Failed to load reference DOCX '%s': %s", reference_docx, e)
                doc = Document()
        else:
            doc = Document()

        # Add metadata as the title page
        has_title_content = any(metadata.get(k) for k in ('title', 'author', 'date'))

        if metadata.get('title'):
            doc.add_heading(metadata.get('title'), 0)

        author_lbl = "Автор" if lang == "ru" else "Author"
        date_lbl = "Дата" if lang == "ru" else "Date"

        if metadata.get('author'):
            doc.add_paragraph(f"{author_lbl}: {metadata.get('author')}")

        if metadata.get('date'):
            doc.add_paragraph(f"{date_lbl}: {metadata.get('date')}")

        if has_title_content:
            doc.add_page_break()

        # Initialize the HTML to DOCX parser
        parser = HtmlToDocx()

        # The parser adds the converted HTML directly into the Document object
        parser.add_html_to_document(html_content, doc)

        # Ensure the output directory exists (guard against empty dirname)
        out_dir = os.path.dirname(output_path)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)

        # Save the Word Document
        doc.save(output_path)
        logger.info("DOCX saved: %s", output_path)
